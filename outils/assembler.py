#!/usr/bin/env python3
"""Assemble le roman : sortie/roman.md, sortie/roman.epub (ebook KDP),
sortie/roman.docx et sortie/roman_broche.pdf (intérieur broché 5,5 × 8,5 po).

Usage : python3 outils/assembler.py
Dépendances : pandoc, python-docx, LibreOffice (soffice), pdftotext, police EB Garamond.
"""
import glob, os, re, shutil, subprocess, tempfile
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt

RACINE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SORTIE = os.path.join(RACINE, "sortie")
TITRE = "Valeur à neuf"
SOUS_TITRE = "roman"
POLICE = "EB Garamond"
LECTEUR = "markdown-subscript-superscript"

CRITIC = re.compile(r"\{(\+\+|--|~~|>>).*?(\+\+|--|~~|<<)\}", re.S)


def pretraiter(texte, nom):
    if nom.startswith("03_"):
        # CriticMarkup affiché tel quel (la légende du chapitre le décrit), hors blocs de code
        blocs = re.split(r"(```.*?```)", texte, flags=re.S)
        texte = "".join(b if b.startswith("```") else
                        CRITIC.sub(lambda m: m.group(0).replace("-", "\\-").replace("~", "\\~"), b)
                        for b in blocs)
    lignes, code = [], False
    for l in texte.split("\n"):
        if l.startswith("```"):
            code = not code
        if not code and l.strip() in ("~", "* * *"):
            signe = "~" if l.strip() == "~" else "\\* \\* \\*"
            lignes += ['::: {.sep custom-style="Separateur"}', signe, ":::"]
        else:
            lignes.append(l)
    return "\n".join(lignes).strip() + "\n"


def assembler_md():
    parties = [f"---\ntitle: \"{TITRE}\"\nsubtitle: \"{SOUS_TITRE}\"\nlang: fr-FR\n"
               f"toc-title: \"Table des matières\"\n---\n"]
    for f in sorted(glob.glob(os.path.join(RACINE, "chapitres", "*.md"))):
        parties.append(pretraiter(open(f, encoding="utf-8").read(), os.path.basename(f)))
    md = "\n\n".join(parties)
    chemin = os.path.join(SORTIE, "roman.md")
    open(chemin, "w", encoding="utf-8").write(md)
    return chemin


CSS = """
body { font-family: serif; line-height: 1.45; }
h1 { text-align: center; font-weight: normal; margin: 3em 0 2em; page-break-before: always; }
p { margin: 0; text-indent: 1.2em; text-align: justify; }
h1 + p, .sep + p, blockquote p, pre + p { text-indent: 0; }
.sep p, div.sep { text-align: center; text-indent: 0; margin: 1em 0; }
pre { white-space: pre-wrap; font-size: 0.72em; margin: 1em 0; }
blockquote { margin: 1em 1.5em; font-size: 0.95em; }
del { text-decoration: line-through; }
.title { text-align: center; margin-top: 30%; font-size: 2em; }
.subtitle { text-align: center; font-style: italic; }
"""


def faire_epub(md):
    css = os.path.join(tempfile.gettempdir(), "roman_epub.css")
    open(css, "w").write(CSS)
    sortie = os.path.join(SORTIE, "roman.epub")
    subprocess.run(["pandoc", "-f", LECTEUR, md, "-o", sortie, "--toc", "--toc-depth=1",
                    "--split-level=1", "--css", css], check=True)
    return sortie


def police(st, nom, taille=None):
    st.font.name = nom
    rpr = st.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.append(rf)
    for att in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rf.get(qn(att)) is not None:
            del rf.attrib[qn(att)]
    for att in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(att), nom)
    if taille:
        st.font.size = Pt(taille)


def style(doc, nom):
    try:
        return doc.styles[nom]
    except KeyError:
        return doc.styles.add_style(nom, 1)


def reference_docx():
    ref = os.path.join(tempfile.gettempdir(), "reference_kdp.docx")
    with open(ref, "wb") as out:
        subprocess.run(["pandoc", "--print-default-data-file", "reference.docx"],
                       stdout=out, check=True)
    doc = Document(ref)
    for s in doc.sections:
        s.page_width, s.page_height = Inches(5.5), Inches(8.5)
        s.top_margin = s.bottom_margin = Inches(0.75)
        s.left_margin, s.right_margin = Inches(0.8), Inches(0.6)
        s.header_distance = s.footer_distance = Inches(0.4)
    for nom in ("Normal", "Body Text", "First Paragraph", "Compact", "Block Text"):
        st = style(doc, nom)
        police(st, POLICE, 11)
        pf = st.paragraph_format
        pf.space_before = pf.space_after = Pt(0)
        pf.line_spacing = 1.2
        pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        pf.first_line_indent = Cm(0.5) if nom == "Body Text" else Cm(0)
        pf.widow_control = True
    bt = style(doc, "Block Text")
    bt.font.size = Pt(10)
    bt.paragraph_format.left_indent = bt.paragraph_format.right_indent = Cm(0.6)
    bt.paragraph_format.space_before = bt.paragraph_format.space_after = Pt(4)
    for nom in ("Source Code", "Verbatim Char"):
        st = style(doc, nom)
        police(st, "Courier New", 7.5)
    sc = style(doc, "Source Code").paragraph_format
    sc.space_before = sc.space_after = Pt(6)
    sc.alignment = WD_ALIGN_PARAGRAPH.LEFT
    h1 = style(doc, "Heading 1")
    police(h1, POLICE, 16)
    h1.font.bold = False
    h1.font.color.rgb = None
    h1.paragraph_format.page_break_before = True
    h1.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    h1.paragraph_format.space_before, h1.paragraph_format.space_after = Cm(3), Cm(1.5)
    for nom, taille in (("Title", 26), ("Subtitle", 13)):
        st = style(doc, nom)
        police(st, POLICE, taille)
        st.font.color.rgb = None
        st.font.bold = False
        st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style(doc, "Title").paragraph_format.space_before = Cm(6)
    style(doc, "Subtitle").font.italic = True
    sep = style(doc, "Separateur")
    sep.base_style = doc.styles["Normal"]
    sep.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sep.paragraph_format.space_before = sep.paragraph_format.space_after = Pt(9)
    sep.paragraph_format.first_line_indent = Cm(0)
    doc.save(ref)
    return ref


def champ_page(paragraphe):
    run = paragraphe.add_run()
    for tag, texte in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), tag)
        else:
            el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = texte
        run._r.append(el)


def mise_en_page(docx):
    doc = Document(docx)
    corps = doc.paragraphs
    # Saut de page après le sous-titre, puis une page blanche (verso de la page de titre)
    premier_titre = next(p for p in corps if p.style.name == "Heading 1")
    blanc = premier_titre.insert_paragraph_before("", style="Normal")
    blanc.paragraph_format.page_break_before = True
    # Le prologue commence sur une page de droite : pas de saut supplémentaire (Heading 1)
    s = doc.sections[0]
    s.different_first_page_header_footer = True
    pied = s.footer.paragraphs[0] if s.footer.paragraphs else s.footer.add_paragraph()
    pied.alignment = WD_ALIGN_PARAGRAPH.CENTER
    champ_page(pied)
    for r in pied.runs:
        r.font.name, r.font.size = POLICE, Pt(9)
    # Marges en miroir (petit fond intérieur)
    reglages = doc.settings.element
    mm = OxmlElement("w:mirrorMargins")
    avant = [reglages.find(qn(t)) for t in ("w:saveFormsData", "w:saveSubsetFonts", "w:embedSystemFonts",
             "w:embedTrueTypeFonts", "w:printFormsData", "w:zoom", "w:view", "w:writeProtection")]
    avant = [e for e in avant if e is not None]
    if avant:
        avant[0].addnext(mm)
    else:
        reglages.insert(0, mm)
    doc.save(docx)


def pdf(docx, dossier):
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", dossier, docx],
                   check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    return os.path.join(dossier, os.path.splitext(os.path.basename(docx))[0] + ".pdf")


def pages_des_titres(pdf_chemin, titres):
    texte = subprocess.run(["pdftotext", "-layout", pdf_chemin, "-"], capture_output=True,
                           text=True, check=True).stdout
    pages = texte.split("\f")
    trouves, depart = {}, 0
    for t in titres:
        cle = re.sub(r"\s+", " ", t).strip()[:25]
        for i in range(depart, len(pages)):
            debut = re.sub(r"\s+", " ", pages[i]).strip()[:120]
            if cle in debut:
                trouves[t], depart = i + 1, i + 1
                break
    return trouves


def ajouter_table(docx, numeros):
    doc = Document(docx)
    titres = [p for p in doc.paragraphs if p.style.name == "Heading 1"]
    p = doc.add_paragraph("Table des matières", style="Heading 1")
    for h in titres:
        e = doc.add_paragraph(style="Normal")
        e.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        e.paragraph_format.space_after = Pt(3)
        e.paragraph_format.tab_stops.add_tab_stop(Inches(4.1), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        e.add_run(f"{h.text}\t{numeros.get(h.text, '')}")
    doc.save(docx)
    return [h.text for h in titres]


def faire_docx(md):
    ref = reference_docx()
    docx = os.path.join(SORTIE, "roman.docx")
    subprocess.run(["pandoc", "-f", LECTEUR, md, "-o", docx, "--reference-doc", ref], check=True)
    mise_en_page(docx)
    titres = [p.text for p in Document(docx).paragraphs if p.style.name == "Heading 1"]
    with tempfile.TemporaryDirectory() as tmp:
        numeros = pages_des_titres(pdf(docx, tmp), titres)
    ajouter_table(docx, numeros)
    with tempfile.TemporaryDirectory() as tmp:
        final = pdf(docx, tmp)
        verif = pages_des_titres(final, titres)
        shutil.copy(final, os.path.join(SORTIE, "roman_broche.pdf"))
    ecarts = {t: (numeros.get(t), verif.get(t)) for t in titres if numeros.get(t) != verif.get(t)}
    return docx, numeros, ecarts


if __name__ == "__main__":
    os.makedirs(SORTIE, exist_ok=True)
    md = assembler_md()
    print("md  :", md)
    print("epub:", faire_epub(md))
    docx, numeros, ecarts = faire_docx(md)
    print("docx:", docx, "| chapitres paginés :", len(numeros), "| écarts :", ecarts or "aucun")
